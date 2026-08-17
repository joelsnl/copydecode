from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from copydecode.document import Segment, default_output_path, detect_format
from copydecode.glossary import Glossary
from copydecode.io import load_document, write_document
from copydecode.router import needs_llm


class FormatDetectTests(unittest.TestCase):
    def test_suffixes(self) -> None:
        self.assertEqual(detect_format(Path("a.epub")), "epub")
        self.assertEqual(detect_format(Path("a.PDF")), "pdf")
        self.assertEqual(detect_format(Path("notes.txt")), "txt")
        self.assertEqual(detect_format(Path("x.markdown")), "md")
        self.assertEqual(detect_format(Path("x.docx")), "docx")
        self.assertEqual(detect_format(Path("x.jsonl")), "jsonl")

    def test_default_output_keeps_type(self) -> None:
        path = Path("book.pdf")
        self.assertEqual(default_output_path(path, "polish").name, "book.polished.pdf")
        self.assertEqual(default_output_path(path, "translate", "txt").name, "book.en.txt")


class TextRoundtripTests(unittest.TestCase):
    def test_txt_paragraphs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "a.txt"
            src.write_text(
                "The mountain wind was cold.\n\nHe could not help but smile.\n",
                encoding="utf-8",
            )
            doc = load_document(src)
            self.assertEqual(len(doc.chapters), 1)
            self.assertEqual(len(doc.chapters[0].segments), 2)
            doc.chapters[0].segments[1].text = "He couldn't help smiling."
            out = Path(tmp) / "a.polished.txt"
            write_document(doc, out)
            text = out.read_text(encoding="utf-8")
            self.assertIn("mountain wind", text)
            self.assertIn("couldn't help smiling", text)

    def test_markdown_keeps_code_and_skips_llm(self) -> None:
        md = "# Title\n\nHe could not help but smile.\n\n```\nkeep_this = True\n```\n"
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "a.md"
            src.write_text(md, encoding="utf-8")
            doc = load_document(src)
            tags = [seg.tag for ch in doc.chapters for seg in ch.segments]
            self.assertIn("pre", tags)
            pre = [seg for ch in doc.chapters for seg in ch.segments if seg.tag == "pre"][0]
            self.assertFalse(needs_llm(pre, "polish", "auto", Glossary()))
            self.assertTrue(
                needs_llm(
                    Segment("x", 0, "He could not help but smile.", "p"),
                    "polish",
                    "auto",
                    Glossary(),
                )
            )
            out = Path(tmp) / "a.out.md"
            write_document(doc, out)
            written = out.read_text(encoding="utf-8")
            self.assertIn("```", written)
            self.assertIn("keep_this = True", written)
            self.assertIn("# Title", written)

    def test_html_roundtrip(self) -> None:
        html = "<html><body><p>The mountain wind was cold.</p><p>He could not help but smile.</p></body></html>"
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "a.html"
            src.write_text(html, encoding="utf-8")
            doc = load_document(src)
            self.assertEqual(len(doc.chapters[0].segments), 2)
            doc.chapters[0].segments[0].text = "The mountain wind cut like a blade."
            out = Path(tmp) / "a.out.html"
            write_document(doc, out)
            data = out.read_text(encoding="utf-8")
            self.assertIn("cut like a blade", data)
            self.assertNotIn("data-np-id", data)

    def test_jsonl_and_cross_format(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "a.jsonl"
            src.write_text(
                json.dumps({"text": "The mountain wind was cold."}) + "\n"
                + json.dumps({"text": "He could not help but smile."}) + "\n",
                encoding="utf-8",
            )
            doc = load_document(src)
            out = Path(tmp) / "a.txt"
            write_document(doc, out, fmt="txt")
            self.assertIn("mountain wind", out.read_text(encoding="utf-8"))


class PdfDocxTests(unittest.TestCase):
    def test_pdf_write_then_load(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "a.txt"
            src.write_text(
                "The mountain wind was cold.\n\nHe could not help but smile.\n",
                encoding="utf-8",
            )
            doc = load_document(src)
            pdf_path = Path(tmp) / "a.pdf"
            write_document(doc, pdf_path, fmt="pdf")
            self.assertGreater(pdf_path.stat().st_size, 200)
            loaded = load_document(pdf_path)
            blob = " ".join(seg.text for ch in loaded.chapters for seg in ch.segments)
            self.assertIn("mountain wind", blob)
            self.assertIn("could not help but smile", blob)

    def test_docx_roundtrip(self) -> None:
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "a.docx"
            d = DocxDocument()
            d.add_heading("Chapter One", level=1)
            d.add_paragraph("The mountain wind was cold.")
            d.add_paragraph("He could not help but smile.")
            d.save(str(src))
            doc = load_document(src)
            texts = [seg.text for ch in doc.chapters for seg in ch.segments]
            self.assertTrue(any("mountain wind" in t for t in texts))
            for ch in doc.chapters:
                for seg in ch.segments:
                    if "could not help" in seg.text:
                        seg.text = "He couldn't help smiling."
            out = Path(tmp) / "a.out.docx"
            write_document(doc, out)
            again = load_document(out)
            blob = " ".join(seg.text for ch in again.chapters for seg in ch.segments)
            self.assertIn("couldn't help smiling", blob)
            self.assertIn("mountain wind", blob)


if __name__ == "__main__":
    unittest.main()
