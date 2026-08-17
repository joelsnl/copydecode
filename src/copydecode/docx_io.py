# Author: joelsnl
"""DOCX paragraph round-trip. Rewritten paragraphs flatten run-level bold/italic."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from copydecode.document import Chapter, Document, Segment, chapter_from_blocks


def _paragraphs(doc: Any) -> list[Any]:
    items: list[Any] = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                items.extend(cell.paragraphs)
    return items


def _heading_tag(style_name: str) -> str | None:
    name = (style_name or "").lower()
    if not name.startswith("heading"):
        return None
    for ch in name:
        if ch.isdigit():
            level = min(6, max(1, int(ch)))
            return f"h{level}"
    return "h1"


def load_docx(path: Path) -> Document:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("DOCX support needs python-docx. Install with: pip install python-docx") from exc
    doc = DocxDocument(str(path))
    paras = _paragraphs(doc)
    chapters: list[Chapter] = []
    current_id = "d0"
    current_title = path.stem
    blocks: list[tuple[str, str, int]] = []

    def flush() -> None:
        nonlocal current_id, current_title, blocks
        if not blocks:
            return
        segs = [
            Segment(item_id=current_id, index=idx, text=text, tag=tag)
            for text, tag, idx in blocks
        ]
        chapters.append(
            Chapter(current_id, f"{path.name}#{current_id}", current_title, None, segs)
        )
        current_id = f"d{len(chapters)}"
        blocks = []

    for index, para in enumerate(paras):
        text = (para.text or "").strip()
        if not text:
            continue
        style_name = para.style.name if para.style is not None else ""
        tag = _heading_tag(style_name) or "p"
        if tag == "h1" and blocks:
            flush()
            current_title = text[:80]
        blocks.append((text, tag, index))
    flush()
    if not chapters:
        chapters.append(chapter_from_blocks("d0", path.name, path.stem, []))
    return Document(
        path=path,
        fmt="docx",
        chapters=chapters,
        payload={"doc": doc, "paragraphs": paras},
        title=path.stem,
    )


def _set_paragraph_text(paragraph: Any, text: str) -> None:
    if (paragraph.text or "") == text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    paragraph.add_run(text)


def write_docx(doc: Document, output: Path, chapters: list[Chapter] | None = None) -> None:
    work = chapters if chapters is not None else doc.chapters
    payload = doc.payload if isinstance(doc.payload, dict) else None
    if payload and payload.get("doc") is not None and doc.fmt == "docx":
        paras = payload["paragraphs"]
        for chapter in work:
            for seg in chapter.segments:
                if 0 <= seg.index < len(paras):
                    _set_paragraph_text(paras[seg.index], seg.text)
        output.parent.mkdir(parents=True, exist_ok=True)
        payload["doc"].save(str(output))
        return
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("DOCX support needs python-docx. Install with: pip install python-docx") from exc
    fresh = DocxDocument()
    if doc.title:
        fresh.add_heading(doc.title, level=1)
    for chapter in work:
        if chapter.skip:
            continue
        if chapter.title and chapter.title != doc.title:
            fresh.add_heading(chapter.title, level=1)
        for seg in chapter.segments:
            if seg.tag.startswith("h") and seg.tag[1:].isdigit():
                fresh.add_heading(seg.text, level=int(seg.tag[1]))
            else:
                fresh.add_paragraph(seg.text)
    output.parent.mkdir(parents=True, exist_ok=True)
    fresh.save(str(output))
